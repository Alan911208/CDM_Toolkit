"""
Test script for cdm_toolkit_plus.py.
Creates test data and runs all new functions.
"""

import os
import sys
import shutil
from pathlib import Path

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from cdm_toolkit_plus import *

OUTPUT_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "test_plus_output")


def setup():
    if os.path.exists(OUTPUT_DIR):
        shutil.rmtree(OUTPUT_DIR)
    os.makedirs(OUTPUT_DIR)
    # Create test folder structure
    test_dir = os.path.join(OUTPUT_DIR, "test_folder")
    os.makedirs(os.path.join(test_dir, "sub_a"))
    os.makedirs(os.path.join(test_dir, "sub_b"))
    Path(os.path.join(test_dir, "file1.xlsx")).touch()
    Path(os.path.join(test_dir, "file2.pdf")).touch()
    Path(os.path.join(test_dir, "sub_a", "file3.txt")).touch()
    Path(os.path.join(test_dir, "sub_b", "file4.xlsx")).touch()
    return test_dir


def test_generate_file_listing():
    print("[TEST] File listing ...", end=" ")
    test_dir = setup()
    out_path = os.path.join(OUTPUT_DIR, "file_listing.xlsx")
    wb = generate_file_listing(test_dir, out_path)
    ws = wb["File Listing"]
    # Should have header + at least 4 files
    assert ws.max_row >= 5, f"Expected >=5 rows, got {ws.max_row}"
    print("OK")


def test_generate_folder_tree():
    print("[TEST] Folder tree ...", end=" ")
    test_dir = setup()
    out_path = os.path.join(OUTPUT_DIR, "folder_tree.xlsx")
    wb = generate_folder_tree(test_dir, out_path)
    assert "Folder Tree" in wb.sheetnames
    print("OK")


def test_track_changes():
    print("[TEST] Track changes (openpyxl fallback) ...", end=" ")
    from openpyxl import Workbook

    # Create original
    orig_path = os.path.join(OUTPUT_DIR, "original.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "Test"
    ws.cell(row=1, column=1, value="Hello")
    ws.cell(row=2, column=1, value="World")
    wb.save(orig_path)

    # Create modified
    wb2 = Workbook()
    ws2 = wb2.active
    ws2.title = "Test"
    ws2.cell(row=1, column=1, value="Hello")
    ws2.cell(row=2, column=1, value="Changed")  # Changed!

    changes = track_changes_with_comments(wb2, orig_path)
    assert changes >= 1, f"Expected >=1 change, got {changes}"
    print(f"OK ({changes} changes)")


if __name__ == "__main__":
    setup()
    tests = [test_generate_file_listing, test_generate_folder_tree, test_track_changes]
    passed = 0
    for t in tests:
        try:
            t()
            passed += 1
        except Exception as e:
            print(f"FAIL: {e}")
    print(f"\nResults: {passed}/{len(tests)} passed")
