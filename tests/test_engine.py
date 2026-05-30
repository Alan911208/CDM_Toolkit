"""Test suite for cdm_engine."""
import os, sys, shutil
from datetime import datetime, timedelta
from openpyxl import Workbook

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from cdm_engine import *

OUTPUT_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "test_output")


def setup():
    """Clean and recreate test_output directory."""
    if os.path.exists(OUTPUT_DIR):
        shutil.rmtree(OUTPUT_DIR)
    os.makedirs(OUTPUT_DIR)


# ================================================================
#  Mock Data Generators
# ================================================================

def make_ae_data(wb):
    """Create AE sheet with 20 rows of realistic data."""
    ws = wb.active
    ws.title = "AE"
    headers = ["STUDYID", "SITEID", "USUBJID", "AESEQ", "AETERM",
               "AESTDTC", "AEENDTC", "AESER", "AESEV", "AEOUT", "AEREL"]
    ws.append(headers)
    apply_header_style(ws)
    subjects = ["001-001"] * 4 + ["001-002"] * 4 + ["001-003"] * 4 + ["001-004"] * 4 + ["001-005"] * 4
    terms = ["Headache"] * 2 + ["Nausea"] * 2 + ["Fatigue"] * 2 + ["Dizziness"] * 2 + ["Rash"] * 2 + \
            ["Headache"] * 2 + ["Nausea"] * 2 + ["Fatigue"] * 2 + ["Dizziness"] * 2 + ["Rash"] * 2
    outcomes = ["Recovered", "Ongoing", "Unknown", "Recovered/Resolved with Sequelae"] * 5
    start_date = datetime(2025, 6, 1)
    for i in range(20):
        ws.append([
            "STUDY001", f"SITE{10 + i % 3:02d}", subjects[i], i + 1, terms[i],
            start_date + timedelta(days=i * 7),
            start_date + timedelta(days=i * 7 + 5),
            "N", "MILD", outcomes[i], "Reasonable Possib",
        ])
    auto_width(ws)
    return ws


def make_lab_data(wb):
    """Create lab data sheet with mixed units."""
    ws = wb.create_sheet("LB")
    headers = ["STUDYID", "USUBJID", "VISIT", "LBTEST", "LBORRES", "LBORRESU"]
    ws.append(headers)
    apply_header_style(ws)
    lab_data = [
        ["STUDY001", "001-001", "Screening", "Glucose", 90, "mg/dL"],
        ["STUDY001", "001-001", "Week 4", "Glucose", 0.9, "g/L"],
        ["STUDY001", "001-003", "Screening", "Hemoglobin", 14.5, "g/dL"],
        ["STUDY001", "001-003", "Week 4", "Hemoglobin", 145, "g/L"],
        ["STUDY001", "001-001", "Screening", "Creatinine", 0.9, "mg/dL"],
        ["STUDY001", "001-003", "Screening", "Creatinine", 88, "umol/L"],
    ]
    for row in lab_data:
        ws.append(row)
    auto_width(ws)
    return ws


# ================================================================
#  Tests
# ================================================================

def test_import():
    """Verify all 18 core functions are callable."""
    funcs = [
        standardise_terms, highlight_duplicates, insert_blank_every_n,
        insert_blank_between_groups, delete_blank_rows, generate_toc,
        generate_sas_derive_toc, split_sheet_by_column, merge_workbooks,
        track_changes, changes_report, add_days_to_column, calc_date_diff,
        convert_lab_units, batch_rename_sheets, scan_special_chars,
        generate_file_listing, generate_folder_tree,
    ]
    for f in funcs:
        assert callable(f), f"{f.__name__} is not callable"
    print(f"[PASS] test_import — {len(funcs)} functions callable")


def test_standardise_terms():
    """Create AE data, run standardise_terms, verify 'Recovered / Resolved' appears."""
    wb = Workbook()
    make_ae_data(wb)
    ws = wb["AE"]

    count = standardise_terms(ws)
    sample = ws.cell(row=2, column=10).value  # AEOUT row 2 was "Recovered"
    assert "Recovered / Resolved" in str(sample), f"Expected 'Recovered / Resolved', got: {sample}"
    wb.save(os.path.join(OUTPUT_DIR, "test_standardise_terms.xlsx"))
    print(f"[PASS] test_standardise_terms ({count} replacements)")


def test_highlight_duplicates():
    """Create AE data, highlight column 5, verify cell fill is yellow."""
    wb = Workbook()
    make_ae_data(wb)
    ws = wb["AE"]

    highlight_duplicates(ws, col_start=5, col_end=5, row_start=2)  # AETERM column
    # Row 3 has "Headache" which duplicates row 2 -> both should be yellow
    fill_rgb = ws.cell(row=3, column=5).fill.start_color.rgb
    assert fill_rgb in ("00FFFF00", "FFFF00"), f"Expected yellow fill, got {fill_rgb}"
    wb.save(os.path.join(OUTPUT_DIR, "test_highlight_duplicates.xlsx"))
    print("[PASS] test_highlight_duplicates")


def test_blank_rows():
    """Create AE data, insert_blank_every_n + insert_blank_between_groups, verify counts > 0."""
    wb = Workbook()
    make_ae_data(wb)
    ws = wb["AE"]

    inserted1 = insert_blank_every_n(ws, n=5, col_letter="A")
    assert inserted1 > 0, f"Expected > 0 blank rows from insert_blank_every_n, got {inserted1}"

    inserted2 = insert_blank_between_groups(ws, col_letter="C")
    assert inserted2 > 0, f"Expected > 0 blank rows from insert_blank_between_groups, got {inserted2}"

    wb.save(os.path.join(OUTPUT_DIR, "test_blank_rows.xlsx"))
    print(f"[PASS] test_blank_rows (every-5: {inserted1}, between-groups: {inserted2})")


def test_generate_toc():
    """Create AE+LB data, generate_toc, verify 'TOC' in sheetnames."""
    wb = Workbook()
    make_ae_data(wb)
    make_lab_data(wb)

    generate_toc(wb, "TOC")
    assert "TOC" in wb.sheetnames, f"TOC not in sheetnames: {wb.sheetnames}"
    wb.save(os.path.join(OUTPUT_DIR, "test_generate_toc.xlsx"))
    print("[PASS] test_generate_toc")


def test_sheet_splitter():
    """Create AE data with severity column, split_sheet_by_column, verify sheets created."""
    wb = Workbook()
    make_ae_data(wb)
    ws = wb["AE"]

    # Add a severity column (col L = 12)
    ws.cell(row=1, column=12, value="SEVERITY_GROUP")
    for i in range(2, ws.max_row + 1):
        ws.cell(row=i, column=12, value="MILD" if i % 3 == 0 else "MODERATE")

    split_sheet_by_column(wb, "AE", split_col="L", header_rows=1)
    assert "MILD" in wb.sheetnames or "MODERATE" in wb.sheetnames, \
        f"Expected MILD or MODERATE sheet, got {wb.sheetnames}"
    wb.save(os.path.join(OUTPUT_DIR, "test_sheet_splitter.xlsx"))
    print("[PASS] test_sheet_splitter")


def test_track_changes():
    """Save original workbook, modify a cell, track_changes, verify >= 1 change detected."""
    orig_path = os.path.join(OUTPUT_DIR, "test_track_changes_original.xlsx")

    wb = Workbook()
    make_ae_data(wb)
    wb.save(orig_path)

    # Modify two cells
    wb["AE"].cell(row=2, column=5).value = "Migraine"  # was "Headache"
    wb["AE"].cell(row=3, column=10).value = "Resolved"  # was "Ongoing"

    changes = track_changes(wb["AE"], orig_path)
    assert changes >= 1, f"Expected >= 1 change, got {changes}"
    wb.save(os.path.join(OUTPUT_DIR, "test_track_changes.xlsx"))
    print(f"[PASS] test_track_changes ({changes} changes)")


def test_date_calc():
    """Create AE data, add_days_to_column + calc_date_diff, verify counts > 0."""
    wb = Workbook()
    make_ae_data(wb)
    ws = wb["AE"]

    # Add 30 days to AE start dates (column F)
    n = add_days_to_column(ws, "F", 30, row_start=2)
    assert n > 0, f"Expected > 0 dates shifted, got {n}"

    # Calculate duration: G - F -> column M
    ws.cell(row=1, column=13, value="AEDUR_DAYS")
    m = calc_date_diff(ws, "F", "G", "M", row_start=2)
    assert m > 0, f"Expected > 0 date diffs computed, got {m}"

    wb.save(os.path.join(OUTPUT_DIR, "test_date_calc.xlsx"))
    print(f"[PASS] test_date_calc (dates shifted: {n}, diffs computed: {m})")


def test_convert_lab_units():
    """Create lab data with mixed units, convert to mg/dL, verify unit changed."""
    wb = Workbook()
    make_lab_data(wb)
    ws = wb["LB"]

    n = convert_lab_units(ws, "D", "E", "F", "mg/dL", test_filter="Glucose")
    assert n > 0, f"Expected > 0 conversions, got {n}"

    # Verify all Glucose rows now have mg/dL unit
    for row_idx in range(2, ws.max_row + 1):
        test_name = ws.cell(row=row_idx, column=4).value
        unit = ws.cell(row=row_idx, column=6).value
        if test_name == "Glucose":
            assert unit == "mg/dL", f"Row {row_idx}: Glucose unit not mg/dL, got {unit}"

    wb.save(os.path.join(OUTPUT_DIR, "test_convert_lab_units.xlsx"))
    print(f"[PASS] test_convert_lab_units (converted: {n})")


def test_full_pipeline():
    """Run all steps (standardise + highlight + convert + date diff + toc) on one workbook."""
    wb = Workbook()
    make_ae_data(wb)
    make_lab_data(wb)

    # 1. Standardise terms
    n1 = standardise_terms(wb["AE"])
    assert n1 > 0, "Expected > 0 term standardisations"

    # 2. Highlight duplicates in AETERM
    highlight_duplicates(wb["AE"], col_start=5, col_end=5, row_start=2)

    # 3. Convert lab units
    n2 = convert_lab_units(wb["LB"], "D", "E", "F", "mg/dL", test_filter="Glucose")
    assert n2 > 0, "Expected > 0 lab conversions"

    # 4. Date difference
    wb["AE"].cell(row=1, column=13, value="AEDUR")
    m = calc_date_diff(wb["AE"], "F", "G", "M", row_start=2)
    assert m > 0, "Expected > 0 date diffs"

    # 5. Generate TOC
    generate_toc(wb, "TOC")
    assert "TOC" in wb.sheetnames, "TOC not generated"

    # Auto-width all sheets
    for ws in wb.worksheets:
        auto_width(ws)

    wb.save(os.path.join(OUTPUT_DIR, "test_full_pipeline.xlsx"))
    print(f"[PASS] test_full_pipeline (terms: {n1}, lab: {n2}, date diffs: {m})")


# ================================================================
#  Main
# ================================================================
if __name__ == "__main__":
    setup()
    tests = [
        test_import, test_standardise_terms, test_highlight_duplicates,
        test_blank_rows, test_generate_toc, test_sheet_splitter,
        test_track_changes, test_date_calc, test_convert_lab_units,
        test_full_pipeline,
    ]
    passed = failed = 0
    for t in tests:
        try:
            t()
            passed += 1
        except Exception as e:
            print(f"[FAIL] {t.__name__} — {e}")
            import traceback; traceback.print_exc()
            failed += 1
    print(f"Results: {passed} PASSED, {failed} FAILED out of {len(tests)}")
